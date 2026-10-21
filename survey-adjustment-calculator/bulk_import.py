"""Bounded XLSX/CSV import with column mapping and an atomic preview.

No spreadsheet expressions are evaluated. All selected point tables are loaded
before observations; errors keep the candidate project from being committed.
"""
from __future__ import annotations

import base64
import csv
import io
import math
import re
import zipfile
from pathlib import Path

import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from adjustment import AdjustmentError, parse_matrix
from network import Network, KINDS, DEFAULT_SIGMA, ANGLES, angle_degrees, example_network, linearize

MAX_FILE = 4*1024*1024
MAX_ROWS = 2000
MAX_COLUMNS = 64
POINT_FIELDS = {'name': ['点号', '点名', 'name', 'point', 'id'],
                'X': ['X', '北坐标', 'northing'], 'Y': ['Y', '东坐标', 'easting'],
                'Z': ['Z'], 'H': ['H', '高程', 'height'], 'fixed': ['固定分量', 'fixed'],
                'fixed_X': ['固定X', 'fixed_x'], 'fixed_Y': ['固定Y', 'fixed_y'],
                'fixed_Z': ['固定Z', 'fixed_z'], 'fixed_H': ['固定H', 'fixed_h'], 'unit': ['坐标单位', '单位', 'unit']}
OBS_FIELDS = {'id': ['编号', '观测编号', 'id'], 'kind': ['类型', '观测类型', 'kind', 'type'],
              'from': ['起点', '测站', 'from', 'station'], 'to': ['终点', '目标点', 'to', 'target'],
              'back': ['后视点', 'back'], 'fore': ['前视点', 'fore'], 'points': ['连接点', 'points'],
              'value': ['观测值', 'value', '实测值'], 'value_x': ['ΔX', 'DX', 'X'],
              'value_y': ['ΔY', 'DY', 'Y'], 'value_z': ['ΔZ', 'DZ', 'Z'], 'value_h': ['H', '高程'],
              'sigma': ['标准差', 'sigma', 'std'], 'sigma_x': ['σX', 'sigma_x', 'std_x'],
              'sigma_y': ['σY', 'sigma_y', 'std_y'], 'sigma_z': ['σZ', 'sigma_z', 'std_z'],
              'sigma_h': ['σH', 'sigma_h', 'std_h'], 'covariance': ['协方差阵', 'covariance'],
              'cxx': ['Cxx'], 'cxy': ['Cxy'], 'cxz': ['Cxz'], 'cyy': ['Cyy'], 'cyz': ['Cyz'], 'czz': ['Czz'],
              'components': ['坐标分量', 'components'], 'enabled': ['参与计算', 'enabled'],
              'unit': ['观测单位', '单位', 'unit'], 'sigma_unit': ['标准差单位', 'sigma_unit'],
              'covariance_unit': ['协方差单位', 'covariance_unit']}


def normalize(value):
    text = str(value).strip().replace('（', '(').replace('）', ')')
    return re.sub(r'[\s_\-]+', '', re.sub(r'\([^)]*\)', '', text)).lower()


def mapping(headers, fields):
    index = {normalize(h): i for i, h in enumerate(headers)}
    return {key: index[next(a for a in map(normalize, aliases) if a in index)]
            for key, aliases in fields.items() if any(a in index for a in map(normalize, aliases))}


def _unit(value):
    return str(value or '').strip().lower().replace('²', '2').replace('^2', '2').replace(' ', '').replace('″', 'arcsec').replace('°', 'deg')


def _length_unit(unit):
    unit = _unit(unit)
    if unit not in ('m', 'mm', 'cm'): raise AdjustmentError('长度单位须为 m、mm 或 cm。')
    return {'m': 1., 'mm': .001, 'cm': .01}[unit]


def _bool(value, default=False):
    if value is None or str(value).strip() == '': return default
    if isinstance(value, bool): return value
    s = str(value).strip().lower()
    if s in ('1', 'true', 'yes', '是', '参与', '固定'): return True
    if s in ('0', 'false', 'no', '否', '停用', '不固定'): return False
    raise AdjustmentError('布尔值请填写 是/否、true/false 或 1/0。')


def _number(value, label):
    if value is None or str(value).strip() == '': raise AdjustmentError(f'请填写{label}。')
    try: result = float(value)
    except (ValueError, TypeError) as exc: raise AdjustmentError(f'{label}须为数值。') from exc
    if not math.isfinite(result): raise AdjustmentError(f'{label}须为有限数。')
    return result


def _vector(value):
    if isinstance(value, str): value = re.split(r'[\s,，;；]+', value.strip())
    return np.asarray(value, dtype=object).reshape(-1).tolist()


def read_tables(filename, content):
    suffix = Path(filename).suffix.lower()
    if suffix not in ('.xlsx', '.csv', '.tsv'):
        raise AdjustmentError('支持 .xlsx、.csv、.tsv；旧 .xls 请先另存为 .xlsx。')
    try: data = base64.b64decode(content, validate=True)
    except (ValueError, TypeError) as exc: raise AdjustmentError('文件内容编码无效。') from exc
    if not data or len(data) > MAX_FILE: raise AdjustmentError('单个导入文件须小于 4 MB，且不能为空。')
    raw_tables = []
    if suffix == '.xlsx':
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                info = archive.infolist()
                if len(info) > 2000 or sum(i.file_size for i in info) > 40*1024*1024:
                    raise AdjustmentError('Excel 解压后数据超过导入上限，请拆分文件。')
                if any(i.filename.endswith('vbaProject.bin') for i in info): raise AdjustmentError('请使用不含宏的 .xlsx 文件。')
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
            try:
                for sheet in book.worksheets:
                    if sheet.max_row > MAX_ROWS+1 or sheet.max_column > MAX_COLUMNS:
                        raise AdjustmentError(f'{sheet.title} 超过 {MAX_ROWS} 行或 {MAX_COLUMNS} 列的导入上限。')
                    rows, formulas = [], []
                    for row in sheet.iter_rows():
                        values = []
                        for cell in row:
                            if cell.data_type == 'f': formulas.append({'row': cell.row, 'column': cell.column, 'message': '公式单元格请先粘贴为数值。'})
                            value = cell.value
                            if value is not None and not isinstance(value, (str, bool, int, float)): value = str(value)
                            if isinstance(value, float) and not math.isfinite(value): raise AdjustmentError(f'{sheet.title}!{cell.coordinate} 含非有限数。')
                            values.append(value)
                        rows.append(values)
                    raw_tables.append((sheet.title, rows, formulas))
            finally: book.close()
        except AdjustmentError: raise
        except Exception as exc: raise AdjustmentError('无法读取 Excel，请确认文件为有效的 .xlsx。') from exc
    else:
        try:
            if data.startswith((b'\xff\xfe', b'\xfe\xff')): text = data.decode('utf-16')
            else:
                try: text = data.decode('utf-8-sig')
                except UnicodeDecodeError: text = data.decode('gb18030')
            delimiter = '\t' if suffix == '.tsv' else None
            if delimiter is None:
                try: delimiter = csv.Sniffer().sniff(text[:8192], delimiters=',;\t').delimiter
                except csv.Error: delimiter = ','
            rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
            if len(rows) > MAX_ROWS+1 or any(len(row) > MAX_COLUMNS for row in rows): raise AdjustmentError('CSV 超过导入行列上限。')
            raw_tables.append((Path(filename).stem, rows, []))
        except AdjustmentError: raise
        except (UnicodeError, csv.Error) as exc: raise AdjustmentError('无法识别 CSV 编码或分隔符。') from exc
    tables = []
    for title, rows, formulas in raw_tables:
        nonempty = [(i, row) for i, row in enumerate(rows) if any(v is not None and str(v).strip() for v in row)]
        if not nonempty: continue
        first, row = nonempty[0]
        headers = [str(value or '').strip() for value in row]
        if len(set(h for h in headers if h)) != len([h for h in headers if h]):
            raise AdjustmentError(f'{title} 有重复列名，请先区分后导入。')
        pm, om = mapping(headers, POINT_FIELDS), mapping(headers, OBS_FIELDS)
        suggested = 'observations' if 'kind' in om or ('from' in om and 'to' in om) else 'points' if 'name' in pm and any(a in pm for a in ('X', 'Y', 'Z', 'H')) else 'ignore'
        tables.append({'name': title, 'filename': str(filename), 'headers': headers,
                       'rows': [r+[None]*(len(headers)-len(r)) for _, r in nonempty[1:]],
                       'source_rows': [i+1 for i, _ in nonempty[1:]], 'formulas': formulas,
                       'type': suggested, 'mapping': om if suggested == 'observations' else pm,
                       'mappings': {'points': pm, 'observations': om}})
    if not tables: raise AdjustmentError('文件中没有可导入的数据表。')
    return {'tables': tables, 'fields': {'points': list(POINT_FIELDS), 'observations': list(OBS_FIELDS)}}


def preview_import(project, tables, *, policy='append', units=None):
    base = Network.from_dict(project)
    if policy not in ('append', 'replace'): raise AdjustmentError('未知导入方式。')
    if not isinstance(tables, list) or not 0 < len(tables) <= 40: raise AdjustmentError('请选择 1 至 40 张数据表。')
    net = Network(base.family, base.name) if policy == 'replace' else base
    net.method = base.method
    units = units or {}
    errors, warnings, added_points, added_observations = [], [], 0, 0
    def error(table, row, message, column=None):
        errors.append({'table': table.get('name', ''), 'file': table.get('filename', ''), 'row': row, 'column': column, 'message': str(message)})
    selected = [t for t in tables if t.get('type') in ('points', 'observations')]
    if not selected: raise AdjustmentError('至少选择一张控制点表或观测表。')
    if sum(len(t.get('rows', [])) for t in selected) > MAX_ROWS: raise AdjustmentError('一次导入最多 2000 行数据。')
    for t in selected:
        if not isinstance(t.get('headers'), list) or not isinstance(t.get('mapping'), dict): raise AdjustmentError('列映射无效。')
        if any(not isinstance(v, int) or isinstance(v, bool) or not 0 <= v < len(t['headers']) for v in t['mapping'].values()):
            raise AdjustmentError('列映射超出表格范围。')
        for item in t.get('formulas', []): error(t, item['row'], item['message'], item['column'])
    for kind in ('points', 'observations'):
        for table in selected:
            if table['type'] != kind: continue
            maps, headers = table['mapping'], table['headers']
            for index, row in enumerate(table['rows']):
                line = table.get('source_rows', [])[index] if index < len(table.get('source_rows', [])) else index+2
                def get(key, default=None):
                    column = maps.get(key)
                    value = row[column] if column is not None and column < len(row) else None
                    return default if value is None or str(value).strip() == '' else value
                def header_unit(key):
                    column = maps.get(key)
                    text = headers[column] if column is not None else ''
                    match = re.search(r'[（(]([^）)]+)[）)]', text)
                    return match.group(1) if match else None
                try:
                    if any(isinstance(v, str) and v.lstrip().startswith('=') for v in row):
                        raise AdjustmentError('公式内容请先粘贴为数值。')
                    if kind == 'points':
                        name = str(get('name', '')).strip()
                        if not name: raise AdjustmentError('请映射并填写点号。')
                        coordinates = {}
                        for axis in ('X', 'Y', 'Z', 'H'):
                            value = get(axis)
                            if value is None and axis in net.axes: raise AdjustmentError(f'缺少 {axis} 坐标或列映射。')
                            if value is not None:
                                scale = _length_unit(get('unit') or header_unit(axis) or units.get('length', 'm'))
                                coordinates[axis] = _number(value, axis)*scale
                        if net.family == '水准网' and get('X') is None and get('Y') is None:
                            coordinates.update(X=(len(net.points)%3)*50., Y=(len(net.points)//3)*60.)
                        raw_fixed = get('fixed', '')
                        if isinstance(raw_fixed, bool) or str(raw_fixed).strip().lower() in ('是', '否', 'true', 'false', '0', '1'):
                            fixed = list(net.axes) if _bool(raw_fixed) else []
                        else:
                            fixed = list(re.sub(r'[\s,，;；]+', '', str(raw_fixed).upper()))
                            if any(axis not in net.axes for axis in fixed): raise AdjustmentError('固定分量须属于当前网型的坐标分量。')
                        for axis in net.axes:
                            if get('fixed_'+axis) is not None:
                                fixed = [k for k in fixed if k != axis]+([axis] if _bool(get('fixed_'+axis)) else [])
                        net.add_point(name, fixed=sorted(set(fixed)), **coordinates)
                        added_points += 1
                    else:
                        raw_kind = str(get('kind', '')).strip()
                        aliases = {normalize(k): k for k in KINDS}
                        aliases.update({normalize(v): k for k, v in KINDS.items()})
                        aliases.update({'gnss': 'gnss', 'gnss基线': 'gnss', '基线': 'gnss'})
                        obs_kind = aliases.get(normalize(raw_kind))
                        if obs_kind is None: raise AdjustmentError('请映射并填写有效的观测类型。')
                        if get('points') is not None: ids = re.split(r'[\s,，;；→]+', str(get('points')).strip())
                        elif obs_kind == 'angle': ids = [get('back'), get('from'), get('fore') or get('to')]
                        elif obs_kind == 'coordinate': ids = [get('from') or get('to')]
                        else: ids = [get('from'), get('to')]
                        if any(v is None for v in ids): raise AdjustmentError('请填写并映射所需端点；水平角需要后视点、测站、前视点。')
                        ids = [str(v).strip() for v in ids]
                        comps = re.split(r'[\s,，;；]+', str(get('components')).strip().upper()) if get('components') else list(net.axes)
                        # XYZ without delimiters is convenient for coordinate tables.
                        if len(comps) == 1 and len(comps[0]) > 1: comps = list(comps[0])
                        vector_axes = ['X', 'Y', 'Z'] if obs_kind == 'gnss' else comps if obs_kind == 'coordinate' else []
                        raw = get('value')
                        if raw is None and vector_axes:
                            components = [get('value_'+a.lower()) for a in vector_axes]
                            if any(v is not None for v in components):
                                if any(v is None for v in components): raise AdjustmentError('观测分量不完整，请填写全部选择的分量。')
                                raw = components
                        if obs_kind in ANGLES:
                            angular_unit = _unit(get('unit') or header_unit('value') or units.get('angle', 'deg'))
                            if angular_unit not in ('deg', 'rad', 'dms'): raise AdjustmentError('角度单位须为 deg、rad 或 dms。')
                            value = None if raw is None else angle_degrees(raw)*(180/math.pi if angular_unit == 'rad' else 1)
                        else:
                            value = None if raw is None else [_number(v, '观测值')*_length_unit(get('unit') or header_unit('value') or units.get('length', 'm')) for v in _vector(raw)]
                            if raw is not None and vector_axes and get('value') is None:
                                value = [_number(get('value_'+a.lower()), '观测分量')*_length_unit(get('unit') or header_unit('value_'+a.lower()) or units.get('length', 'm')) for a in vector_axes]
                        raw_sigma = get('sigma')
                        sigma_components = [get('sigma_'+a.lower()) for a in vector_axes]
                        if raw_sigma is None and any(v is not None for v in sigma_components):
                            if any(v is None for v in sigma_components): raise AdjustmentError('分量标准差不完整。')
                            raw_sigma = sigma_components
                        sigma_unit = _unit(get('sigma_unit') or header_unit('sigma') or (units.get('angle_sigma', 'arcsec') if obs_kind in ANGLES else units.get('length_sigma', 'm')))
                        if obs_kind in ANGLES:
                            if sigma_unit not in ('arcsec', 'deg', 'rad'): raise AdjustmentError('角度标准差单位须为 arcsec、deg 或 rad。')
                            sigma_scale = {'arcsec': 1., 'deg': 3600., 'rad': 180*3600/math.pi}[sigma_unit]
                        else: sigma_scale = _length_unit(sigma_unit)
                        sigma = DEFAULT_SIGMA[obs_kind] if raw_sigma is None else [_number(v, '标准差')*sigma_scale for v in _vector(raw_sigma)]
                        if raw_sigma is not None and get('sigma') is None and vector_axes and obs_kind not in ANGLES:
                            sigma = [_number(get('sigma_'+a.lower()), '分量标准差')*_length_unit(get('sigma_unit') or header_unit('sigma_'+a.lower()) or units.get('length_sigma', 'm')) for a in vector_axes]
                        cov = get('covariance')
                        if cov is not None: cov = parse_matrix(str(cov), '协方差阵')
                        elif obs_kind == 'gnss' and any(get(k) is not None for k in ('cxx', 'cxy', 'cxz', 'cyy', 'cyz', 'czz')):
                            cxx, cyy, czz = [_number(get(k), k) for k in ('cxx', 'cyy', 'czz')]
                            cxy, cxz, cyz = [_number(get(k, 0), k) for k in ('cxy', 'cxz', 'cyz')]
                            cov = np.array([[cxx, cxy, cxz], [cxy, cyy, cyz], [cxz, cyz, czz]])
                        if cov is not None:
                            cu = _unit(get('covariance_unit') or header_unit('covariance') or units.get('covariance', 'rad2' if obs_kind in ANGLES else 'm2'))
                            scales = {'rad2': 1., 'deg2': (math.pi/180)**2, 'arcsec2': (math.pi/(180*3600))**2} if obs_kind in ANGLES else {'m2': 1., 'cm2': .0001, 'mm2': .000001}
                            if cu not in scales: raise AdjustmentError('协方差单位不适用于该观测。')
                            cov = (cov*scales[cu]).tolist()
                        elif raw_sigma is None:
                            warnings.append(f'{table["name"]} 第 {line} 行未填写标准差，采用 {DEFAULT_SIGMA[obs_kind]}'+('″' if obs_kind in ANGLES else ' m')+'。')
                        identity = get('id')
                        net.add_observation(obs_kind, ids, value=value, sigma=sigma, covariance=cov, components=comps, name=str(identity) if identity is not None else None)
                        net.observations[-1]['enabled'] = _bool(get('enabled'), True)
                        added_observations += 1
                except (AdjustmentError, ValueError, TypeError) as exc: error(table, line, exc)
    statistics = None
    if not errors:
        try: statistics = linearize(net).stats
        except AdjustmentError as exc: error({'name': '模型校验'}, None, exc)
    return {'ok': not errors, 'project': net.to_dict() if not errors else None,
            'errors': errors[:200], 'error_count': len(errors), 'warnings': warnings[:100],
            'added_points': added_points, 'added_observations': added_observations, 'statistics': statistics,
            'policy': policy}


def template(family, kind='xlsx'):
    net = example_network(family, measured=True)
    point_headers = ['点号', 'X(m)', 'Y(m)', 'Z(m)', 'H(m)', '固定分量']
    point_rows = [[name, p['X'], p['Y'], p['Z'], p['H'], ','.join(p['fixed'])] for name, p in net.points.items()]
    obs_headers = ['编号', '类型', '起点', '终点', '后视点', '前视点', '观测值', '标准差', '坐标分量', '协方差阵', '参与计算']
    obs_rows = []
    for o in net.observations:
        ids = o['points']; angle = o['kind'] == 'angle'
        text = lambda value: ' '.join(str(v) for v in value) if isinstance(value, list) else value
        cov = '\n'.join(' '.join(str(v) for v in row) for row in o['covariance']) if o['covariance'] else None
        obs_rows.append([o['id'], KINDS[o['kind']], ids[1] if angle else ids[0], ids[-1] if len(ids)>1 else None,
                         ids[0] if angle else None, ids[2] if angle else None, text(o['value']), text(o['sigma']), ','.join(o['components']), cov, True])
    if kind in ('points', 'observations'):
        output = io.StringIO(); writer = csv.writer(output)
        writer.writerow(point_headers if kind == 'points' else obs_headers)
        writer.writerows(point_rows if kind == 'points' else obs_rows)
        data = output.getvalue().encode('utf-8-sig')
        filename = family+('-控制点.csv' if kind == 'points' else '-观测.csv')
    else:
        book = Workbook(); book.remove(book.active)
        for title, headers, rows in [('控制点', point_headers, point_rows), ('观测', obs_headers, obs_rows)]:
            sheet = book.create_sheet(title); sheet.append(headers)
            for row in rows: sheet.append(row)
            sheet.freeze_panes = 'B2'; sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]: cell.font = Font(name='Microsoft YaHei', bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='5453D7'); cell.alignment = Alignment(vertical='center')
            sheet.row_dimensions[1].height = 26
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = 18 if column[0].value != '协方差阵' else 44
                for cell in list(column)[1:]: cell.font = Font(name='Microsoft YaHei', size=11); cell.alignment = Alignment(vertical='top', wrap_text=True)
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if isinstance(cell.value, (float, int)) and not isinstance(cell.value, bool): cell.number_format = '0.##########'
        info = book.create_sheet('导入说明')
        for line in ['这些是可计算的示例数据，正式观测请替换。', '点号须唯一；长度与长度标准差为米。',
                     '角度值为十进制度或 123:45:56.7，角度标准差为角秒。', 'GNSS 协方差为同一 XYZ 框架中的 3×3 对称正定阵，单位 m²。',
                     '固定分量用 X,Y,Z 或 H；参与计算用 true/false。', '导入前可预览并调整列映射；错误不会部分写入项目。',
                     '仅导入数值，不计算 Excel 公式；公式结果请先粘贴为值。']:
            info.append([line])
        info.column_dimensions['A'].width = 100
        for row in info: row[0].font = Font(name='Microsoft YaHei', size=11); row[0].alignment = Alignment(wrap_text=True, vertical='top')
        output = io.BytesIO(); book.save(output); data = output.getvalue(); filename = family+'-批量导入模板.xlsx'
    return {'filename': filename, 'content_base64': base64.b64encode(data).decode('ascii')}
